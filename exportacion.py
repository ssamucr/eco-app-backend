"""Exportacion de la vista de analitica a Excel y PDF. Ambos formatos salen de las mismas tablas."""
import io
from datetime import date, datetime
from typing import Optional

from fastapi.responses import Response

from analitica import DIAS_SEMANA, construir_analitica
from main import SALDO_INICIAL, _cuentas_detalle, _rechazar, app, fetch_all

SECCIONES = ["resumen", "ciclos", "gastos", "ritmo", "patrones", "cuentas", "metas", "deuda", "obligaciones", "movimientos"]
TITULOS = {
    "resumen": "Resumen", "ciclos": "Ciclos", "gastos": "Gastos", "ritmo": "Ritmo", "patrones": "Patrones",
    "cuentas": "Cuentas", "metas": "Metas", "deuda": "Deuda", "obligaciones": "Obligaciones", "movimientos": "Movimientos",
}
TIPO_MOVIMIENTO = {"GASTO": "Gasto", "INGRESO": "Ingreso", "TRANSFERENCIA": "Transferencia", "PAGO_TARJETA": "Pago a tarjeta"}
TIPO_CUENTA = {"AHORRO": "Ahorros", "CORRIENTE": "Corriente", "TARJETA_CREDITO": "Tarjeta de crédito", "EFECTIVO": "Efectivo"}
NIVEL = {"alerta": "Atención", "info": "Dato", "ok": "Bien"}


def _col(titulo, tipo="texto", peso=None):
    return {"titulo": titulo, "tipo": tipo, "peso": peso}


def _rango(c):
    return "%s al %s" % (c["fecha_inicio"].strftime("%d/%m/%Y"), c["fecha_fin"].strftime("%d/%m/%Y"))


def _movimientos(ciclo):
    filas = fetch_all(
        "SELECT t.fecha, t.tipo, c.nombre AS categoria, t.descripcion, o.nombre AS origen, d.nombre AS destino, "
        "t.monto, t.referencia FROM transacciones t LEFT JOIN categorias c ON c.id_categoria = t.id_categoria "
        "LEFT JOIN cuentas o ON o.id_cuenta = t.id_cuenta_origen LEFT JOIN cuentas d ON d.id_cuenta = t.id_cuenta_destino "
        "WHERE t.fecha >= ? AND t.fecha <= ? ORDER BY t.fecha DESC, t.id_transaccion DESC LIMIT 5000",
        (ciclo["fecha_inicio"], ciclo["fecha_fin"]),
    )
    return filas


def _tablas(d, claves):
    """Lista de secciones: cada una con sus tablas {titulo, columnas, filas}. Una celda puede ser (valor, tipo)."""
    ciclo, k, s = d["ciclo"], d["kpis"], d["situacion"]
    salida = []

    def seccion(clave, *tablas):
        if clave in claves:
            salida.append({"clave": clave, "titulo": TITULOS[clave], "tablas": [t for t in tablas if t["filas"]]})

    def tabla(titulo, columnas, filas):
        return {"titulo": titulo, "columnas": columnas, "filas": filas}

    ritmo = d["ritmo"]
    indicadores = [
        ["Período", _rango(ciclo)],
        ["Configuración de ciclo", ciclo["config"]],
        ["Ingresos del ciclo", (k["ingresos"], "dinero")],
        ["Gastos del ciclo", (k["gastos"], "dinero")],
        ["Balance del ciclo", (k["balance"], "dinero")],
        ["Tasa de ahorro", (k["tasa_ahorro"], "pct")],
        ["Ingresos vs. ciclo anterior", (k["ingresos_variacion_pct"], "pct")],
        ["Gastos vs. ciclo anterior", (k["gastos_variacion_pct"], "pct")],
        ["Ahorro total", (s["ahorro"], "dinero")],
        ["Deuda en tarjetas", (s["deuda"], "dinero")],
        ["Patrimonio neto", (s["patrimonio_neto"], "dinero")],
        ["Gasto diario promedio", (s["gasto_diario"], "dinero")],
        ["Días de gasto que cubre tu ahorro", (s["cobertura_dias"], "entero")],
        ["Proyección de gasto al cierre del ciclo", (ritmo["proyeccion_cierre"] if ritmo["proyeccion_confiable"] else None, "dinero")],
    ]
    seccion(
        "resumen",
        # las filas sin valor (por ejemplo, sin base para comparar) no se muestran
        tabla("Indicadores", [_col("Indicador"), _col("Valor")], [f for f in indicadores if _valor(f[1], "texto")[0] is not None]),
        tabla("Hallazgos", [_col("Tipo", peso=9), _col("Detalle", peso=30), _col("Cómo se calcula", peso=52)],
              [[NIVEL[i["nivel"]], i["texto"], "\n".join(i["explicacion"])] for i in d["insights"]]),
    )
    seccion(
        "ciclos",
        tabla("Ingresos y gastos por ciclo", [
            _col("Inicio", "fecha"), _col("Fin", "fecha"), _col("Días", "entero"), _col("Ingresos", "dinero"),
            _col("Gastos", "dinero"), _col("Balance", "dinero"), _col("Tasa de ahorro", "pct"),
            _col("Ahorro al cierre", "dinero"), _col("Deuda al cierre", "dinero"), _col("Patrimonio al cierre", "dinero"),
        ], [[h["fecha_inicio"], h["fecha_fin"], h["dias"], h["ingresos"], h["gastos"], h["balance"], h["tasa_ahorro"],
             h["patrimonio"]["ahorro"], h["patrimonio"]["deuda"], h["patrimonio"]["neto"]] for h in d["historial"]]),
    )
    g = d["gastos"]
    seccion(
        "gastos",
        tabla("Gastos por categoría", [
            _col("Categoría"), _col("Total", "dinero"), _col("% del gasto", "pct"), _col("Ciclo anterior", "dinero"),
            _col("Variación vs. anterior", "pct"), _col("Promedio por ciclo", "dinero"), _col("Variación vs. promedio", "pct"),
        ], [[c["nombre"], c["total"], c["porcentaje"], c["anterior"], c["variacion_pct"], c["promedio"], c["vs_promedio_pct"]]
            for c in g["categorias"]]),
        tabla("Mayores gastos", [_col("Fecha", "fecha"), _col("Descripción"), _col("Categoría"), _col("Cuenta"), _col("Monto", "dinero")],
              [[m["fecha"], m["descripcion"], m["categoria"] or "Sin categoría", m["cuenta"], m["monto"]] for m in g["mayores"]]),
        tabla("Gasto por cuenta o tarjeta", [_col("Cuenta"), _col("Tipo"), _col("Total", "dinero"), _col("% del gasto", "pct")],
              [[c["nombre"], "Tarjeta" if c["es_tarjeta"] else "Cuenta", c["total"], c["porcentaje"]] for c in g["por_cuenta"]]),
        tabla("Gastos hormiga", [_col("Umbral", "dinero"), _col("Cantidad", "entero"), _col("Total", "dinero"), _col("% del gasto", "pct")],
              [[g["hormiga"]["umbral"], g["hormiga"]["cantidad"], g["hormiga"]["total"], g["hormiga"]["porcentaje"]]]),
        tabla("Gastos hormiga: detalle", [_col("Descripción"), _col("Veces", "entero"), _col("Total", "dinero")],
              [[h["descripcion"], h["veces"], h["total"]] for h in g["hormiga"]["detalle"]]),
    )
    seccion(
        "ritmo",
        tabla("Gasto acumulado por día del ciclo", [
            _col("Día del ciclo", "entero"), _col("Ciclo actual", "dinero"), _col("Ciclo anterior", "dinero"), _col("Proyección", "dinero"),
        ], [[p["dia"], p["actual"], p["anterior"], p["proyeccion"]] for p in ritmo["puntos"]]),
    )
    p = d["patrones"]
    seccion(
        "patrones",
        tabla("Gasto por día de la semana", [_col("Día"), _col("Total", "dinero"), _col("Gastos", "entero"), _col("Promedio por gasto", "dinero")],
              [[x["dia"], x["total"], x["cantidad"], x["promedio"]] for x in p["por_dia_semana"]]),
        tabla("Gasto promedio por día del ciclo", [_col("Día del ciclo", "entero"), _col("Promedio", "dinero")],
              [[x["dia"], x["promedio"]] for x in p["por_dia_ciclo"]]),
        tabla("Gastos recurrentes", [
            _col("Descripción"), _col("Categoría"), _col("Veces", "entero"), _col("Ciclos", "entero"), _col("Promedio", "dinero"),
            _col("Último", "dinero"), _col("Por ciclo", "dinero"), _col("Última vez", "fecha"),
        ], [[r["descripcion"], r["categoria"], r["veces"], r["ciclos"], r["promedio"], r["ultimo"], r["por_ciclo"], r["ultima_fecha"]]
            for r in p["recurrentes"]]),
    )
    cuentas = _cuentas_detalle()
    seccion(
        "cuentas",
        tabla("Cuentas y tarjetas", [_col("Cuenta"), _col("Entidad"), _col("Tipo"), _col("Saldo", "dinero"), _col("Límite", "dinero"), _col("Uso", "pct")],
              [[c["nombre"], c["entidad"], TIPO_CUENTA.get(c["tipo"], c["tipo"]), c["saldo"], c.get("limite_credito"), c.get("uso_pct")] for c in cuentas]),
        tabla("Subcuentas", [_col("Cuenta"), _col("Subcuenta"), _col("Saldo", "dinero"), _col("Meta", "dinero")],
              [[c["nombre"], sc["nombre"], sc["saldo"], sc["saldo_meta"]] for c in cuentas for sc in c.get("subcuentas", [])]),
    )
    seccion(
        "metas",
        tabla("Avance de metas", [
            _col("Subcuenta"), _col("Cuenta"), _col("Saldo", "dinero"), _col("Meta", "dinero"), _col("Avance", "pct"),
            _col("Falta", "dinero"), _col("Aporte del plan", "dinero"), _col("Aportes libres (est.)", "dinero"), _col("Ciclos restantes", "entero"),
        ], [[m["nombre"], m["cuenta"], m["saldo"], m["meta"], m["porcentaje"], m["falta"], m["aporte_plan"], m["aporte_libre"], m["ciclos_restantes"]]
            for m in d["metas"]]),
    )
    de = d["deuda"]
    seccion(
        "deuda",
        tabla("Tarjetas de crédito", [_col("Tarjeta"), _col("Deuda", "dinero"), _col("Límite", "dinero"), _col("Uso", "pct"), _col("Disponible", "dinero")],
              [[t["nombre"], t["saldo"], t["limite"], t["uso_pct"], t["disponible"]] for t in de["tarjetas"]]),
        tabla("Próximas cuotas de financiamientos", [
            _col("Financiamiento"), _col("Cuota"), _col("Vencimiento", "fecha"), _col("Monto", "dinero"), _col("Días para vencer", "entero"),
        ], [[q["descripcion"], "%d de %d" % (q["numero_cuota"], q["numero_cuotas"]), q["fecha_vencimiento"], q["monto"], q["dias_para_vencer"]]
            for q in de["cuotas_proximas"]]),
    )
    ob = d["obligaciones"]
    seccion(
        "obligaciones",
        tabla("Por persona", [_col("Persona"), _col("Te deben", "dinero"), _col("Debes", "dinero"), _col("Neto", "dinero"), _col("Días de la más antigua", "entero")],
              [[x["nombre"], x["te_deben"], x["debes"], x["neto"], x["dias_mas_antigua"]] for x in ob["personas"]]),
        tabla("Antigüedad de lo pendiente", [_col("Rango"), _col("Te deben", "dinero"), _col("Debes", "dinero"), _col("Reposiciones", "dinero")],
              [[x["rango"], x["te_deben"], x["debes"], x["reposiciones"]] for x in ob["antiguedad"]]),
    )
    if "movimientos" in claves:
        seccion(
            "movimientos",
            tabla("Movimientos del período (%s)" % _rango(ciclo), [
                _col("Fecha", "fecha"), _col("Tipo"), _col("Categoría"), _col("Descripción"), _col("Cuenta origen"),
                _col("Cuenta destino"), _col("Monto", "dinero"), _col("Referencia"),
            ], [[m["fecha"], TIPO_MOVIMIENTO.get(m["tipo"], m["tipo"]), m["categoria"], m["descripcion"], m["origen"], m["destino"],
                 m["monto"], m["referencia"]] for m in _movimientos(ciclo)]),
        )
    return salida


# ---------- formatos de salida ----------

def _valor(celda, tipo):
    if isinstance(celda, tuple):
        return celda
    return celda, tipo


def _texto(valor, tipo):
    if valor is None or valor == "":
        return ""
    if tipo == "dinero":
        return ("-" if valor < 0 else "") + "$" + format(abs(float(valor)), ",.2f")
    if tipo == "pct":
        return "%.1f%%" % float(valor)
    if tipo == "fecha":
        return valor.strftime("%d/%m/%Y") if isinstance(valor, (date, datetime)) else str(valor)
    if tipo == "entero":
        return str(int(valor))
    return str(valor)


def _xlsx(d, secciones):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    formatos = {"dinero": '"$"#,##0.00;[Red]-"$"#,##0.00', "pct": '0.0"%"', "fecha": "dd/mm/yyyy", "entero": "0"}
    libro = Workbook()
    libro.remove(libro.active)
    encabezado = PatternFill("solid", fgColor="2B2B2B")
    for sec in secciones:
        hoja = libro.create_sheet(sec["titulo"][:31])
        hoja["A1"] = "Eco · Analítica · %s · %s" % (sec["titulo"], _rango(d["ciclo"]))
        hoja["A1"].font = Font(bold=True, size=13)
        fila = 3
        anchos = {}
        for t in sec["tablas"]:
            hoja.cell(row=fila, column=1, value=t["titulo"]).font = Font(bold=True, size=11)
            fila += 1
            for j, c in enumerate(t["columnas"], start=1):
                celda = hoja.cell(row=fila, column=j, value=c["titulo"])
                celda.font, celda.fill = Font(bold=True, color="FFFFFF"), encabezado
                celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                anchos[j] = max(anchos.get(j, 0), len(c["titulo"]))
            inicio_datos = fila + 1
            for datos in t["filas"]:
                fila += 1
                for j, (c, bruto) in enumerate(zip(t["columnas"], datos), start=1):
                    valor, tipo = _valor(bruto, c["tipo"])
                    celda = hoja.cell(row=fila, column=j, value=valor)
                    if tipo in formatos and valor is not None:
                        celda.number_format = formatos[tipo]
                    if tipo == "texto":
                        celda.alignment = Alignment(vertical="top", wrap_text=True)
                    anchos[j] = max(anchos.get(j, 0), min(len(_texto(valor, tipo)), 60))
            if len(sec["tablas"]) == 1 and len(t["filas"]) > 1:
                hoja.auto_filter.ref = "A%d:%s%d" % (inicio_datos - 1, get_column_letter(len(t["columnas"])), fila)
                hoja.freeze_panes = "A%d" % inicio_datos
            fila += 2
        for j, ancho in anchos.items():
            hoja.column_dimensions[get_column_letter(j)].width = min(max(ancho + 3, 11), 62)
    if not libro.sheetnames:
        libro.create_sheet("Analítica")
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _pdf(d, secciones):
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=estilos["Title"], fontName="Helvetica-Bold", fontSize=18, alignment=0, spaceAfter=2)
    sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=9, textColor=colors.HexColor("#666666"), spaceAfter=10)
    h1 = ParagraphStyle("h1", parent=estilos["Heading2"], fontName="Helvetica-Bold", fontSize=13, spaceBefore=12, spaceAfter=4, keepWithNext=1)
    h2 = ParagraphStyle("h2", parent=estilos["Heading3"], fontName="Helvetica-Bold", fontSize=10, spaceBefore=8, spaceAfter=3, keepWithNext=1)
    celda_txt = ParagraphStyle("c", parent=estilos["Normal"], fontSize=8, leading=10)
    celda_num = ParagraphStyle("n", parent=celda_txt, alignment=2)
    cabecera = ParagraphStyle("hd", parent=celda_txt, textColor=colors.white, fontName="Helvetica-Bold")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title="Eco - Analítica", author="Eco")
    ancho = landscape(A4)[0] - 28 * mm

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawString(14 * mm, 8 * mm, "Eco · Analítica · %s" % _rango(d["ciclo"]))
        canvas.drawRightString(landscape(A4)[0] - 14 * mm, 8 * mm, "Página %d" % documento.page)
        canvas.restoreState()

    historia = [Paragraph("Eco · Analítica", titulo),
                Paragraph("%s · ciclo «%s» · generado el %s" % (_rango(d["ciclo"]), d["ciclo"]["config"], date.today().strftime("%d/%m/%Y")), sub)]

    def grafica_ciclos():
        h = d["historial"]
        dibujo = Drawing(ancho, 150)
        barras = VerticalBarChart()
        barras.x, barras.y, barras.width, barras.height = 40, 25, ancho - 70, 105
        barras.data = [[x["ingresos"] for x in h], [x["gastos"] for x in h]]
        barras.categoryAxis.categoryNames = [x["fecha_inicio"].strftime("%d/%m") for x in h]
        barras.bars[0].fillColor, barras.bars[1].fillColor = colors.HexColor("#1baf7a"), colors.HexColor("#eb6834")
        barras.bars[0].strokeColor = barras.bars[1].strokeColor = None
        barras.barSpacing, barras.groupSpacing = 2, 14
        barras.valueAxis.valueMin = 0
        barras.categoryAxis.labels.fontSize = barras.valueAxis.labels.fontSize = 7
        dibujo.add(barras)
        dibujo.add(String(40, 138, "Ingresos (verde) y gastos (naranja) por ciclo", fontSize=8, fillColor=colors.HexColor("#666666")))
        return dibujo

    for sec in secciones:
        if sec["clave"] == "ciclos" and len(d["historial"]) > 1:
            historia.append(KeepTogether([Paragraph(sec["titulo"], h1), grafica_ciclos()]))
        else:
            historia.append(Paragraph(sec["titulo"], h1))
        for t in sec["tablas"]:
            historia.append(Paragraph(t["titulo"], h2))
            columnas = t["columnas"]
            datos = [[Paragraph(c["titulo"], cabecera) for c in columnas]]
            for fila in t["filas"]:
                renglon = []
                for c, bruto in zip(columnas, fila):
                    valor, tipo = _valor(bruto, c["tipo"])
                    texto = _texto(valor, tipo).replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br/>")
                    renglon.append(Paragraph(texto, celda_txt if tipo in ("texto", "fecha") else celda_num))
                datos.append(renglon)
            peso = [c["peso"] or (max(len(c["titulo"]), 8) if c["tipo"] != "texto" else 28) for c in columnas]
            # una tabla angosta no se estira a todo el ancho: cada punto de peso mide ~3 mm
            total = min(ancho, sum(peso) * 3.3 * mm)
            anchos = [total * p / sum(peso) for p in peso]
            tabla = Table(datos, colWidths=anchos, repeatRows=1, hAlign="LEFT")
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2b2b2b")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f2ee")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#dddddd")),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            historia.append(tabla)
        historia.append(Spacer(1, 4))
    if len(historia) == 2:
        historia.append(Paragraph("No hay datos para las secciones elegidas.", celda_txt))
    doc.build(historia, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


@app.get("/analitica/exportar")
def exportar_analitica(
    formato: str = "xlsx", id_ciclo: Optional[int] = None, ciclos: int = 6, secciones: Optional[str] = None
):
    """Descarga la analitica del ciclo en Excel o PDF. `secciones` (separadas por coma) limita lo que se incluye."""
    if formato not in ("xlsx", "pdf"):
        raise _rechazar(422, "El formato debe ser xlsx o pdf.")
    claves = [s for s in (secciones or "").split(",") if s] or list(SECCIONES)
    desconocidas = [s for s in claves if s not in SECCIONES]
    if desconocidas:
        raise _rechazar(422, "Sección desconocida: %s." % ", ".join(desconocidas))
    datos = construir_analitica(id_ciclo, ciclos)
    # se respeta el orden de la vista, no el del parametro
    lista = _tablas(datos, [s for s in SECCIONES if s in claves])
    if formato == "xlsx":
        contenido = _xlsx(datos, lista)
        tipo = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        contenido = _pdf(datos, lista)
        tipo = "application/pdf"
    nombre = "eco-analitica-%s-%s.%s" % (datos["ciclo"]["fecha_inicio"].isoformat(), datos["ciclo"]["fecha_fin"].isoformat(), formato)
    return Response(content=contenido, media_type=tipo, headers={"Content-Disposition": 'attachment; filename="%s"' % nombre})
